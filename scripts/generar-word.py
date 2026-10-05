"""Genera la guía breve de entrega. Requiere python-docx, solo para este documento."""
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'entrega' / 'TP5-Network-QoS-Guia.docx'
OUTPUT.parent.mkdir(exist_ok=True)
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin = section.bottom_margin = Cm(1.8)
section.left_margin = section.right_margin = Cm(2)
normal = doc.styles['Normal']
normal.font.name, normal.font.size = 'Calibri', Pt(10.5)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.05
for name in ('Title', 'Heading 1', 'Heading 2'):
    doc.styles[name].font.color.rgb = RGBColor.from_string('2454A6')
doc.styles['Heading 1'].font.size = Pt(16)
doc.styles['Heading 2'].font.size = Pt(12)
doc.core_properties.title = 'TP5 — Network QoS Monitor | Instalación y uso'
doc.core_properties.author = ''  # El usuario completa su identidad en la portada.
footer = section.footer.paragraphs[0]
footer.alignment = 2
footer.add_run('Network QoS Monitor · ')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

def p(text, style=None):
    doc.add_paragraph(text, style)

def heading(text):
    doc.add_heading(text, level=1)

def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Shading Accent 1'
    for cell, text in zip(t.rows[0].cells, headers):
        cell.text = text
    for row in rows:
        for cell, text in zip(t.add_row().cells, row):
            cell.text = text
    # Evita cortar una fila entre páginas al editar o imprimir el documento.
    for row in t.rows:
        prop = row._tr.get_or_add_trPr()
        prop.append(OxmlElement('w:cantSplit'))

doc.add_heading('TP5 — Network QoS Monitor', 0)
p('Guía breve · Explicación, instalación y uso')
table(['Datos de la entrega', 'Completar'], [
    ('Institución', '________________________________________'),
    ('Carrera', '________________________________________'),
    ('Asignatura / comisión', '________________________________________'),
    ('Nombre y apellido', '________________________________________'),
    ('Legajo', '________________________________________'),
    ('Docente', '________________________________________'),
    ('Fecha de entrega', '________________________________________'),
    ('Enlace del repositorio', '________________________________________'),
    ('Enlace o nombre del video', '________________________________________'),
])
heading('1. Qué hace el sistema')
p('Aplicación Android para observar la calidad de una conexión. Identifica la red '
  'y el operador disponible; mide latencia TCP/UDP y velocidad de descarga/subida; '
  'guarda resultados con fecha y ubicación cuando existe una posición válida. '
  'Permite consultar gráficos, mapa de calor, filtros y exportar CSV/JSON.')
p('El monitoreo periódico repite lotes de latencia, guarda sus resultados y emite '
  'alertas según umbrales configurables. Se inicia y detiene manualmente. '
  'No ejecuta pruebas automáticas de velocidad.')
doc.add_heading('Componentes', level=2)
p('Interfaz en React Native y TypeScript; módulos Kotlin para mediciones, GPS, '
  'SQLite y servicio Android; estado compartido con Context + useReducer. '
  'Backend propio Node.js/Express para transferencias y eco UDP. Mapa '
  'OpenStreetMap con Leaflet. Entrega Android, probada en Samsung A55 con Android 16.')
p('El APK funciona sin Metro. El servidor de la PC solo es necesario para las '
  'pruebas de velocidad y eco UDP contra ese servidor. El código y la documentación '
  'técnica ampliada se encuentran en el repositorio.')

doc.add_page_break()
heading('2. Instalación y preparación')
doc.add_heading('En el teléfono', level=2)
for text in [
    'Copiar NetworkQoSMonitor-Android.apk al teléfono Android ARM64. Es un APK '
    'de pruebas (Android 7 o posterior); el dispositivo validado es el Samsung A55.',
    'Abrir el archivo y, si Android lo solicita, permitir la instalación desde '
    'esa aplicación de archivos. Instalar y abrir Network QoS Monitor.',
    'Conceder los permisos solicitados para las funciones utilizadas: ubicación '
    '(preferentemente precisa), notificaciones y telefonía cuando se solicite. '
    'Activar también Ubicación en el sistema.',
    'En Monitoreo, tocar Habilitar ubicación antes de medir. Los registros sin '
    'posición reciente se conservan, pero no aparecen como puntos del mapa.',
]:
    p(text, 'List Number')
doc.add_heading('Servidor propio en una PC Windows', level=2)
p('Requiere Node.js 22.11 o posterior y el código del repositorio. PC y teléfono '
  'deben estar en la misma red Wi-Fi para utilizar la dirección privada de la PC.')
p('Desde PowerShell, dentro de la carpeta del proyecto:')
for line in ['cd backend', 'npm.cmd ci', 'cd ..',
             '.\\scripts\\start-backend.ps1 -BindAddress <IP_DE_LA_PC>']:
    run = doc.add_paragraph().add_run(line)
    run.font.name, run.font.size = 'Consolas', Pt(9)
p('Reemplazar <IP_DE_LA_PC> por la dirección IPv4 real de la PC, consultable con '
  'ipconfig. Dejar abierta la terminal del servidor. Si el firewall bloquea la '
  'conexión, permitir TCP 5050 y UDP 5051 en la red privada.')
p('Comprobar desde Chrome del teléfono: http://<IP_DE_LA_PC>:5050/health. '
  'Debe aparecer la identificación network-qos. En Velocidad usar '
  'http://<IP_DE_LA_PC>:5050; para eco UDP usar esa IP y puerto 5051.')
doc.add_heading('Compilar desde el código (opcional)', level=2)
p('El README explica cómo preparar JDK/SDK portables, aceptar licencias e instalar '
  'dependencias con npm.cmd ci dentro de mobile. Después, desde la raíz: '
  '.\\scripts\\build-android.ps1 -Mode Standalone. '
  'Salida: mobile/android/app/build/outputs/apk/release/app-release.apk. '
  'No hace falta Android Studio ni emulador. El backend incluye un Dockerfile '
  'opcional e instrucciones propias; el contenedor no fue probado en esta PC.')

doc.add_page_break()
heading('3. Uso del sistema')
table(['Sección', 'Operación'], [
    ('Inicio', 'Consultar conexión activa, operador y señal disponible.'),
    ('Latencia', 'Configurar al menos tres destinos. Ejemplo: PC:5051 UDP, '
     '8.8.8.8:53 TCP y 9.9.9.9:53 TCP. Iniciar y esperar 10 intentos por destino. '
     'Los DNS públicos no son servidores de eco UDP.'),
    ('Velocidad', 'Indicar URL del backend y tamaño 1, 5 o 10 MiB. Iniciar las '
     'tres rondas de descarga/subida; para una demostración breve, usar 1 MiB.'),
    ('Monitoreo', 'Usa los destinos guardados por la última prueba de Latencia. '
     'Iniciar, comprobar Activo y detener desde la app o su notificación. '
     'Detener antes de realizar mediciones manuales.'),
    ('Historial', 'Actualizar, tocar una fecha para abrir su gráfico debajo. '
     'Filtrar por red, fechas AAAA-MM-DD o área sur,oeste,norte,este. '
     'Exportar CSV/JSON con los filtros aplicados.'),
    ('Mapa', 'Actualizar y observar las muestras de RTT con ubicación. '
     'Tocar un punto para leer RTT y precisión. El calor representa demora, '
     'no cobertura celular; densidad y zoom también afectan el color.'),
])
doc.add_heading('Cómo interpretar los números', level=2)
p('RTT (ms): demora de ida y vuelta. Mínimo/promedio/máximo usan respuestas '
  'válidas. Jitter (ms): variación entre RTT consecutivos; menor suele ser más '
  'estable. Mbps: tasa de transferencia; MiB: tamaño de archivo. '
  '10/10 intentos no significa diez respuestas. Un guion indica dato insuficiente. '
  'Los fallos TCP o ecos UDP ausentes no identifican por sí solos la causa del problema.')
p('Valores iniciales de monitoreo: pausa de 60 s después de cada lote, umbral '
  'de RTT 500 ms, jitter 100 ms y fallos 50 %. Son configurables y de prueba. '
  'Las alertas se separan al menos cinco minutos. Android puede limitar o diferir '
  'el servicio; no continúa tras Forzar detención ni se inicia solo al encender.')
doc.add_heading('Alcance y comprobaciones', level=2)
p('Se verificaron TCP/UDP, velocidad, dos sesiones con RTT y ubicación, mapa, '
  'gráficos, exportaciones reales CSV/JSON y generación/detención del monitoreo. '
  'La condición de pantalla bloqueada queda pendiente de confirmación explícita. '
  'El mensaje de búsqueda de ubicación es informativo y no confirma un GPS actual.')
p('La IP privada de la PC no es accesible por datos móviles. En 4G/5G se pueden '
  'medir destinos TCP públicos; velocidad y UDP requieren un servidor accesible '
  'desde Internet. Una transferencia contra la PC mide el trayecto local, no '
  'necesariamente la velocidad contratada. Las exportaciones pueden contener '
  'coordenadas personales. La entrega incluye video de 3–5 minutos con dos '
  'sesiones reales; no se exige que sean de redes distintas.')
doc.save(OUTPUT)
print(OUTPUT)
